"""Tạo file luyện tập cho Thực hành 6 và 7 (toàn bộ dữ liệu giả lập).

    python tools/practice_files.py

Kết quả nằm trong app/static/labs/ để học viên tải từ bài học. Cùng seed cho cùng một file,
nên số liệu chữa bài trong giáo trình và đáp án giảng viên luôn khớp. Cần openpyxl.
"""
import random
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parent.parent / "app" / "static" / "labs"


def bao_cao_thang():
    """Thực hành 6: BaoCao_Thang.xlsx – sheet VanHanh ẩn, cột GhiChuNoiBo ẩn."""
    random.seed(2026)
    vung_ds = ["Phường Bắc", "Phường Nam", "Xã Đông", "Xã Tây"]
    ho = ["Nguyễn", "Trần", "Lê", "Phạm", "Hoàng", "Võ", "Đặng", "Bùi"]
    dem = ["Văn", "Thị", "Minh", "Thu", "Đức", "Ngọc"]
    ten = ["An", "Bình", "Chi", "Dũng", "Hà", "Khoa", "Lan", "Nam", "Phúc", "Trang"]

    wb = Workbook()
    th = wb.active
    th.title = "TongHop"
    ct = wb.create_sheet("ChiTiet")
    vh = wb.create_sheet("VanHanh")
    ct.append(["MaKH", "TenKH", "DienThoai", "DiaChi", "Vung", "kWh_T10", "SoTienNo", "SoNgayQuaHan", "GhiChuNoiBo"])
    tong = {v: [0, 0] for v in vung_ds}
    for i in range(1, 401):
        vung = random.choice(vung_ds)
        kwh = random.randint(60, 900)
        no = random.choice([0, 0, 0, random.randint(150, 2500) * 1000])
        qua_han = random.randint(5, 60) if no else 0
        ct.append([f"KH{i:05d}", f"{random.choice(ho)} {random.choice(dem)} {random.choice(ten)}", f"0900{i:06d}",
                   f"Số {random.randint(1, 200)}, tổ {random.randint(1, 30)}, {vung}", vung, kwh, no, qua_han,
                   random.choice(["", "", "Hay khiếu nại", "Đã hẹn trả", "Hộ khó khăn"])])
        tong[vung][0] += 1
        tong[vung][1] += kwh
    ct.column_dimensions["I"].hidden = True
    th.append(["Vung", "SoHo", "SanLuong_kWh"])
    for v in vung_ds:
        th.append([v, *tong[v]])
    vh.append(["Tram", "Vung", "TaiCaoNhat_%", "SoLanMatDien_T10"])
    for k, v in enumerate(vung_ds, 1):
        vh.append([f"TBA-{k:02d}", v, random.randint(55, 98), random.randint(0, 6)])
    vh.sheet_state = "hidden"
    wb.save(OUT / "BaoCao_Thang.xlsx")
    no_count = sum(1 for r in ct.iter_rows(min_row=2, values_only=True) if r[6])
    print("BaoCao_Thang.xlsx:", ct.max_row - 1, "hộ;", no_count, "hộ còn nợ;", tong)


def _khong_dau(s):
    s = unicodedata.normalize("NFD", s.replace("Đ", "D").replace("đ", "d"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower().replace(" ", "")


GHI_CHU_TH7 = [
    "Hay vắng nhà, liên hệ con gái {nt} – {sdt}",
    "Chủ hộ đang chạy thận định kỳ, cần báo trước khi cắt điện",
    "Hẹn trả nợ ngày 15, đã gọi nhắc 2 lần",
    "Đang khiếu nại chỉ số công tơ tháng 7",
    "Nhà có người khuyết tật, ưu tiên khôi phục điện",
    "Thuê nhà, chủ nhà là {nt} – {sdt}",
    "Đã nhắc nợ qua Zalo, khách phản ứng gay gắt",
    "Người cao tuổi sống một mình",
]


def san_luong_q3():
    """Thực hành 7: DLX_SanLuong_Q3.xlsx – Cảng Mới chỉ có một xưởng; một hộ An Phú dùng điện rất lớn.

    Sản lượng, số nợ, tên và địa chỉ lấy từ random.seed(7) như bản CSV trước, nên mọi số chữa bài giữ nguyên.
    Các cột thông tin cá nhân thêm vào (CCCD, ngày sinh, giới tính, email, hộ chính sách, tài khoản, ghi chú)
    dùng một bộ sinh số riêng để không làm lệch các số đó.
    """
    random.seed(7)
    phu = random.Random(77)
    phuong_ds = {"An Phú": 160, "Bình Minh": 140, "Đông Hòa": 120, "Tân Lập": 150, "Phú Thượng": 110, "Cảng Mới": 0}
    ho = ["Nguyễn", "Trần", "Lê", "Phạm", "Hoàng", "Võ", "Đặng", "Bùi"]
    ten = ["An", "Bình", "Chi", "Dũng", "Hà", "Khoa", "Lan", "Nam", "Phúc", "Trang"]
    khach = []

    def them(phuong, loai, kwh_tb):
        i = len(khach) + 1
        ho_kh, ten_kh = random.choice(ho), random.choice(ten)
        dia_chi = f"Số {random.randint(1, 300)}, {phuong}"
        nu = phu.random() < 0.5
        sinh = date(phu.randint(1950, 2000), phu.randint(1, 12), phu.randint(1, 28))
        the_ky = (2 if sinh.year >= 2000 else 0) + (1 if nu else 0)
        ghi_chu = ""
        if phu.random() < 0.12:
            ghi_chu = phu.choice(GHI_CHU_TH7).format(nt=f"{phu.choice(ho)} Thị {phu.choice(ten)}",
                                                    sdt=f"09{phu.randint(10, 99)}{phu.randint(0, 999999):06d}")
        khach.append({
            "ma_kh": f"PE{i:06d}", "ho_ten": f"{ho_kh} {'Thị' if nu else 'Văn'} {ten_kh}",
            "so_cccd": f"0{phu.choice(['01', '31', '79', '92'])}{the_ky}{sinh.year % 100:02d}{phu.randint(0, 999999):06d}",
            "ngay_sinh": sinh, "gioi_tinh": "Nữ" if nu else "Nam",
            "dien_thoai": f"0900{i:06d}",
            "email": f"{_khong_dau(ho_kh + ten_kh)}{phu.randint(10, 99)}@gmail.com" if phu.random() < 0.6 else "",
            "dia_chi": dia_chi, "phuong": phuong, "loai_kh": loai, "kwh_tb": kwh_tb,
            "ho_chinh_sach": (phu.choices(["", "Hộ nghèo", "Hộ cận nghèo"], [90, 5, 5])[0] if loai == "Sinh hoạt" else ""),
            "so_tk": (f"{phu.choice(['Vietcombank', 'BIDV', 'Agribank', 'VietinBank'])} {phu.randint(10 ** 9, 10 ** 10 - 1)}"
                      if phu.random() < 0.3 else ""),
            "ghi_chu": ghi_chu})

    for p, n in phuong_ds.items():
        for _ in range(n):
            them(p, "Sinh hoạt", random.randint(120, 450))
        for _ in range(random.randint(8, 15) if n else 0):
            them(p, "Kinh doanh", random.randint(600, 2500))
    them("An Phú", "Sinh hoạt", 3200)
    them("Cảng Mới", "Sản xuất", 42000)

    cot = ["ma_kh", "ho_ten", "so_cccd", "ngay_sinh", "gioi_tinh", "dien_thoai", "email", "dia_chi", "phuong", "loai_kh",
           "thang", "kwh", "so_tien_no", "ho_chinh_sach", "so_tk_ngan_hang", "ghi_chu"]
    wb = Workbook()
    ws = wb.active
    ws.title = "DuLieu"
    ws.append(cot)
    for k in khach:
        for thang in ("2026-07", "2026-08", "2026-09"):
            kwh = round(k["kwh_tb"] * random.uniform(0.85, 1.15))
            no = random.choice([0, 0, 0, 0, round(kwh * 2.5) * 1000])
            ws.append([k["ma_kh"], k["ho_ten"], k["so_cccd"], k["ngay_sinh"], k["gioi_tinh"], k["dien_thoai"], k["email"],
                       k["dia_chi"], k["phuong"], k["loai_kh"], thang, kwh, no, k["ho_chinh_sach"], k["so_tk"], k["ghi_chu"]])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 4).number_format = "dd/mm/yyyy"
        ws.cell(r, 13).number_format = "#,##0"
        for c in (3, 6):                       # CCCD, điện thoại: giữ dạng chữ để không mất số 0 đầu
            ws.cell(r, c).number_format = "@"
    bang = Table(displayName="SanLuong", ref=f"A1:P{ws.max_row}")
    bang.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(bang)
    ws.freeze_panes = "A2"
    for c, w in zip("ABCDEFGHIJKLMNOP", [11, 20, 15, 12, 9, 13, 26, 22, 12, 12, 9, 8, 12, 15, 26, 52]):
        ws.column_dimensions[c].width = w

    mt = wb.create_sheet("MoTa")
    mt.append(["Cột", "Nội dung"])
    for row in [
        ("ma_kh", "Mã khách hàng trên hệ thống thu tiền điện"),
        ("ho_ten", "Họ tên chủ hợp đồng"),
        ("so_cccd", "Số căn cước công dân của chủ hợp đồng"),
        ("ngay_sinh", "Ngày sinh của chủ hợp đồng"),
        ("gioi_tinh", "Giới tính của chủ hợp đồng"),
        ("dien_thoai", "Số điện thoại liên hệ"),
        ("email", "Email nhận hóa đơn điện tử (để trống nếu không đăng ký)"),
        ("dia_chi", "Địa chỉ nơi dùng điện"),
        ("phuong", "Phường nơi dùng điện"),
        ("loai_kh", "Sinh hoạt / Kinh doanh / Sản xuất"),
        ("thang", "Tháng ghi chỉ số (năm-tháng)"),
        ("kwh", "Sản lượng điện trong tháng (kWh)"),
        ("so_tien_no", "Số tiền còn nợ tại thời điểm xuất (đồng)"),
        ("ho_chinh_sach", "Hộ nghèo / Hộ cận nghèo – dùng để áp giá hỗ trợ"),
        ("so_tk_ngan_hang", "Tài khoản đăng ký trích nợ tự động"),
        ("ghi_chu", "Ghi chú của nhân viên thu ngân, chăm sóc khách hàng (gõ tự do)"),
    ]:
        mt.append(row)
    mt.append([])
    mt.append(["Lưu ý", "Dữ liệu giả lập cho Thực hành 7, không phải khách hàng thật. File đóng vai kết quả chức năng “Xuất danh sách” – đủ mọi cột."])
    mt.column_dimensions["A"].width = 18
    mt.column_dimensions["B"].width = 90
    wb.save(OUT / "DLX_SanLuong_Q3.xlsx")
    print("DLX_SanLuong_Q3.xlsx:", len(khach), "khách hàng,", ws.max_row - 1, "dòng,", len(cot), "cột")


if __name__ == "__main__":
    bao_cao_thang()
    san_luong_q3()
